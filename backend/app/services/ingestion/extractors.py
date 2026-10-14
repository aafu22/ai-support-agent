"""Validate and extract text from uploads. Raises ValueError with a user-safe message."""
import io
import zipfile

from app.core.config import settings

Page = tuple[int | None, str]

ALLOWED_EXTENSIONS = {"pdf", "docx", "txt"}
MAX_DOCX_UNCOMPRESSED = 50 * 1024 * 1024  # guards against zip bombs


def detect_type(filename: str, data: bytes) -> str:
    """Check extension AND file signature; never trust the extension alone."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError("Unsupported file type. Upload a PDF, DOCX or TXT file.")
    if not data:
        raise ValueError("The file is empty.")
    if ext == "pdf" and not data.startswith(b"%PDF-"):
        raise ValueError("This file is not a valid PDF.")
    if ext == "docx":
        if not data.startswith(b"PK\x03\x04"):
            raise ValueError("This file is not a valid DOCX.")
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if "word/document.xml" not in z.namelist():
                    raise ValueError("This file is not a valid DOCX.")
                if sum(i.file_size for i in z.infolist()) > MAX_DOCX_UNCOMPRESSED:
                    raise ValueError("The DOCX file is too large once unpacked.")
        except zipfile.BadZipFile:
            raise ValueError("This file is not a valid DOCX.")
    if ext == "txt" and b"\x00" in data[:4096]:
        raise ValueError("This file does not look like plain text.")
    return ext


def extract_pdf(data: bytes) -> list[Page]:
    import pymupdf  # PyMuPDF

    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception:
        raise ValueError("Could not open the PDF (corrupted or encrypted?).")
    with doc:
        if doc.needs_pass:
            raise ValueError("Password-protected PDFs are not supported.")
        if doc.page_count > settings.max_pdf_pages:
            raise ValueError(f"PDF has too many pages (limit {settings.max_pdf_pages}).")
        return [(i + 1, page.get_text("text")) for i, page in enumerate(doc)]

def extract_docx(data: bytes) -> list[Page]:
    from docx import Document as DocxDocument

    try:
        doc = DocxDocument(io.BytesIO(data))
    except Exception:
        raise ValueError("Could not read the DOCX file.")
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return [(None, "\n\n".join(parts))]


def extract_txt(data: bytes) -> list[Page]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    return [(None, text)]


def extract(file_type: str, data: bytes) -> list[Page]:
    return {"pdf": extract_pdf, "docx": extract_docx, "txt": extract_txt}[file_type](data)
